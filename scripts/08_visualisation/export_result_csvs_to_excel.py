#!/usr/bin/env python3
"""Export result-folder CSV files into documented multi-sheet Excel workbooks."""

from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORTS = (
    ("results/tables/amr", "amr_results_tables.xlsx"),
    ("results/tables/mlst_amr_phylogeny", "mlst_amr_phylogeny_results_tables.xlsx"),
)
INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")


# Resolve the repository root from this script or the project override.
def project_root() -> Path:
    override = os.environ.get("EFAECALIS_PROJECT_ROOT")
    return Path(override).resolve() if override else Path(__file__).resolve().parents[2]


# Create a valid, deterministic, and unique Excel worksheet name.
def worksheet_name(stem: str, used: set[str]) -> str:
    base = INVALID_SHEET_CHARACTERS.sub("_", stem)[:31] or "Sheet"
    candidate = base
    counter = 2
    while candidate.casefold() in used:
        suffix = f"_{counter}"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.casefold())
    return candidate


# Convert unambiguous numeric and Boolean fields while preserving identifiers as text.
def excel_value(value: str):
    if value == "":
        return None
    if value in {"TRUE", "FALSE"}:
        return value == "TRUE"
    if re.fullmatch(r"-?(?:0|[1-9]\d*)", value):
        try:
            return int(value)
        except ValueError:
            return value
    if re.fullmatch(r"-?(?:\d+\.\d*|\d*\.\d+)(?:[eE][+-]?\d+)?|-?\d+[eE][+-]?\d+", value):
        try:
            return float(value)
        except ValueError:
            return value
    return value


# Apply concise spreadsheet formatting without changing source values.
def format_worksheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        values = [sheet.cell(row=row, column=column).value for row in range(1, min(sheet.max_row, 300) + 1)]
        width = min(50, max(10, max((len(str(value)) for value in values if value is not None), default=0) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width


# Import one CSV exactly once and report its workbook cardinality.
def add_csv_sheet(workbook: Workbook, csv_path: Path, sheet_name: str) -> tuple[int, int]:
    sheet = workbook.create_sheet(sheet_name)
    with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        rows = 0
        columns = 0
        for row_index, row in enumerate(reader, start=1):
            rows += 1
            columns = max(columns, len(row))
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index, value=value if row_index == 1 else excel_value(value))
    format_worksheet(sheet)
    return max(0, rows - 1), columns


# Build one workbook per requested results folder with a source-to-sheet manifest.
def export_folder(root: Path, relative_folder: str, output_filename: str, overwrite: bool) -> Path:
    folder = root / relative_folder
    csv_paths = sorted(folder.glob("*.csv"), key=lambda path: path.name.casefold())
    if not csv_paths:
        raise RuntimeError(f"No CSV files found in {folder}")
    output_path = folder / output_filename
    if output_path.exists() and not overwrite:
        raise RuntimeError(f"Refusing to overwrite existing workbook: {output_path}")

    workbook = Workbook()
    manifest = workbook.active
    manifest.title = "Manifest"
    manifest.append(["Source CSV", "Worksheet", "Data rows", "Columns"])
    used = {"manifest"}
    for csv_path in csv_paths:
        sheet_name = worksheet_name(csv_path.stem, used)
        data_rows, columns = add_csv_sheet(workbook, csv_path, sheet_name)
        manifest.append([csv_path.name, sheet_name, data_rows, columns])
    format_worksheet(manifest)
    workbook.save(output_path)
    return output_path


# Export both result collections only after validating command-line intent.
def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    root = project_root()
    for relative_folder, output_filename in EXPORTS:
        output = export_folder(root, relative_folder, output_filename, args.overwrite)
        print(f"Created {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
